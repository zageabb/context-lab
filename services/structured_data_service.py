from __future__ import annotations

import csv
import json
import re
import sqlite3
from pathlib import Path

from openpyxl import load_workbook

from models import StructuredDataSource, StructuredDataTable
from services.file_storage import ensure_environment_directories


def structured_db_path(base_data_dir: Path, environment_id: int) -> Path:
    base_dir = ensure_environment_directories(base_data_dir, environment_id)
    return base_dir / "structured_data" / "environment_structured_data.db"


def import_structured_file_to_sqlite(
    base_data_dir: Path,
    environment_id: int,
    source: StructuredDataSource,
) -> list[StructuredDataTable]:
    db_path = structured_db_path(base_data_dir, environment_id)
    db_path.parent.mkdir(parents=True, exist_ok=True)
    suffix = Path(source.original_filename).suffix.lower()
    table_specs = _read_source_tables(Path(source.file_path), suffix)
    imported_tables: list[StructuredDataTable] = []

    with sqlite3.connect(db_path) as connection:
        for index, spec in enumerate(table_specs, start=1):
            sqlite_table_name = _sqlite_safe_table_name(environment_id, source.id, index, spec["display_name"])
            _replace_table(connection, sqlite_table_name, spec["columns"], spec["rows"])
            imported_tables.append(
                StructuredDataTable(
                    source=source,
                    display_name=spec["display_name"],
                    sqlite_table_name=sqlite_table_name,
                    row_count=len(spec["rows"]),
                    column_count=len(spec["columns"]),
                    column_names_json=json.dumps(spec["columns"]),
                )
            )
    return imported_tables


def delete_structured_tables(base_data_dir: Path, environment_id: int, tables: list[StructuredDataTable]) -> None:
    db_path = structured_db_path(base_data_dir, environment_id)
    if not db_path.exists() or not tables:
        return
    with sqlite3.connect(db_path) as connection:
        for table in tables:
            connection.execute(f'DROP TABLE IF EXISTS "{table.sqlite_table_name}"')
        connection.commit()


def preview_structured_table(base_data_dir: Path, environment_id: int, table: StructuredDataTable, limit: int = 25) -> dict:
    db_path = structured_db_path(base_data_dir, environment_id)
    columns = json.loads(table.column_names_json or "[]")
    if not db_path.exists():
        return {"columns": columns, "rows": []}
    with sqlite3.connect(db_path) as connection:
        cursor = connection.execute(f'SELECT * FROM "{table.sqlite_table_name}" LIMIT ?', (limit,))
        rows = [list(row) for row in cursor.fetchall()]
    return {"columns": columns, "rows": rows}


def execute_read_only_sql(base_data_dir: Path, environment_id: int, sql: str, row_limit: int = 200) -> dict:
    query = (sql or "").strip()
    if not query:
        raise ValueError("Enter a SQL query to run.")
    if not _is_read_only_sql(query):
        raise ValueError("Only single read-only SELECT or WITH queries are allowed in the playground.")

    db_path = structured_db_path(base_data_dir, environment_id)
    if not db_path.exists():
        return {"columns": [], "rows": [], "row_count": 0}

    with sqlite3.connect(db_path) as connection:
        connection.row_factory = sqlite3.Row
        cursor = connection.execute(query)
        rows = cursor.fetchmany(max(1, min(row_limit, 1000)))
        columns = [column[0] for column in (cursor.description or [])]
    return {
        "columns": columns,
        "rows": [list(row) for row in rows],
        "row_count": len(rows),
    }


def summarize_structured_tables(tables: list[StructuredDataTable]) -> str:
    if not tables:
        return "No structured data tables are available for this environment."
    lines = []
    for table in tables[:20]:
        columns = json.loads(table.column_names_json or "[]")
        preview = ", ".join(columns[:6]) if columns else "no columns recorded"
        lines.append(
            f"- {table.display_name}: {table.row_count} rows, {table.column_count} columns ({preview})"
        )
    return "\n".join(lines)


def retrieve_structured_rows(
    base_data_dir: Path,
    environment_id: int,
    tables: list[StructuredDataTable],
    query: str,
    row_limit: int = 6,
) -> list[dict]:
    db_path = structured_db_path(base_data_dir, environment_id)
    if not db_path.exists() or not tables:
        return []
    query_terms = _query_terms(query)
    if not query_terms:
        return []

    matches: list[dict] = []
    with sqlite3.connect(db_path) as connection:
        for table in tables:
            columns = json.loads(table.column_names_json or "[]")
            table_terms = _query_terms(f"{table.display_name} {' '.join(columns)}")
            cursor = connection.execute(f'SELECT * FROM "{table.sqlite_table_name}" LIMIT 200')
            for row in cursor.fetchall():
                row_values = ["" if value is None else str(value) for value in row]
                combined = " | ".join(row_values).lower()
                row_terms = _query_terms(combined)
                value_hits = sum(1 for term in query_terms if term in combined)
                metadata_hits = sum(1 for term in query_terms if term in table_terms)
                score = value_hits + (metadata_hits * 1.5)
                if score <= 0:
                    continue
                matches.append(
                    {
                        "table_name": table.display_name,
                        "score": score,
                        "columns": columns,
                        "row_values": row_values,
                        "matched_on": {
                            "value_terms": [term for term in query_terms if term in row_terms],
                            "metadata_terms": [term for term in query_terms if term in table_terms],
                        },
                    }
                )
    matches.sort(key=lambda item: item["score"], reverse=True)
    return matches[:row_limit]


def format_structured_row_context(rows: list[dict]) -> str:
    if not rows:
        return "No structured-data rows were matched for this question."
    sections = []
    for row in rows:
        paired = [f"{column}: {value}" for column, value in zip(row["columns"], row["row_values"])]
        sections.append(
            f"Table: {row['table_name']} | Score: {row['score']}\n" + "\n".join(paired)
        )
    return "\n\n---\n\n".join(sections)


def _read_source_tables(path: Path, suffix: str) -> list[dict]:
    if suffix == ".csv":
        return [_csv_spec(path)]
    if suffix == ".xlsx":
        return _xlsx_specs(path)
    raise ValueError(f"Unsupported structured-data file type: {suffix}")


def _csv_spec(path: Path) -> dict:
    with path.open("r", encoding="utf-8", errors="ignore", newline="") as handle:
        rows = list(csv.reader(handle))
    return _tabular_spec(path.stem, rows)


def _xlsx_specs(path: Path) -> list[dict]:
    workbook = load_workbook(filename=str(path), data_only=True)
    specs: list[dict] = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(["" if value is None else str(value) for value in row])
        spec = _tabular_spec(sheet.title, rows)
        if spec["columns"]:
            specs.append(spec)
    return specs


def _tabular_spec(name: str, rows: list[list[str]]) -> dict:
    non_empty_rows = [list(row) for row in rows if any(str(cell).strip() for cell in row)]
    if not non_empty_rows:
        return {"display_name": name, "columns": [], "rows": []}
    header = _unique_column_names(non_empty_rows[0])
    normalized_rows = []
    for row in non_empty_rows[1:]:
        normalized_rows.append(_pad_row(row, len(header)))
    return {"display_name": name, "columns": header, "rows": normalized_rows}


def _unique_column_names(raw_columns: list[str]) -> list[str]:
    names: list[str] = []
    seen: dict[str, int] = {}
    for index, raw in enumerate(raw_columns, start=1):
        base = _column_safe_name(raw or f"column_{index}")
        count = seen.get(base, 0)
        seen[base] = count + 1
        names.append(base if count == 0 else f"{base}_{count + 1}")
    return names


def _column_safe_name(value: str) -> str:
    cleaned = re.sub(r"[^a-zA-Z0-9]+", "_", value.strip().lower()).strip("_")
    return cleaned or "column"


def _sqlite_safe_table_name(environment_id: int, source_id: int, sheet_index: int, display_name: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", display_name.strip().lower()).strip("_") or "table"
    return f"env_{environment_id}_src_{source_id}_{sheet_index}_{slug}"


def _replace_table(connection: sqlite3.Connection, table_name: str, columns: list[str], rows: list[list[str]]) -> None:
    connection.execute(f'DROP TABLE IF EXISTS "{table_name}"')
    if not columns:
        connection.commit()
        return
    column_sql = ", ".join(f'"{column}" TEXT' for column in columns)
    connection.execute(f'CREATE TABLE "{table_name}" ({column_sql})')
    if rows:
        placeholders = ", ".join("?" for _ in columns)
        connection.executemany(
            f'INSERT INTO "{table_name}" VALUES ({placeholders})',
            [_pad_row(row, len(columns)) for row in rows],
        )
    connection.commit()


def _pad_row(row: list[str], width: int) -> list[str]:
    values = ["" if value is None else str(value) for value in row[:width]]
    return values + [""] * (width - len(values))


def _query_terms(query: str) -> list[str]:
    return [term for term in re.findall(r"[a-zA-Z0-9]{3,}", (query or "").lower())[:12]]


def _is_read_only_sql(query: str) -> bool:
    stripped = query.strip().rstrip(";").strip()
    if not stripped:
        return False
    if ";" in stripped:
        return False
    normalized = stripped.lower()
    if not (normalized.startswith("select") or normalized.startswith("with")):
        return False
    blocked_terms = (
        "insert",
        "update",
        "delete",
        "drop",
        "alter",
        "create",
        "replace",
        "attach",
        "detach",
        "pragma",
        "vacuum",
        "reindex",
        "truncate",
    )
    return not any(re.search(rf"\b{term}\b", normalized) for term in blocked_terms)
